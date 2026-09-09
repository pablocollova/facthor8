#!/usr/bin/env python3
"""Create the editable Facthor8 brand manual."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BRAND = ROOT / "brand"
PNG = BRAND / "png"
OUT = BRAND / "manual"
OUT.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT / "Facthor8_Manual_de_Marca.docx"

NAVY = "06294D"
DARK = "041C34"
GREEN = "00D39A"
ICE = "F7FAFC"
SLATE = "7890A5"
LIGHT_BORDER = "D7E0E7"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=140, start=160, bottom=140, end=160) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LIGHT_BORDER, size="8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def remove_paragraph_border(element) -> None:
    p_pr = element.find(qn("w:pPr"))
    if p_pr is None:
        return
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_run_font(run, name="DejaVu Sans", size=10.5, bold=False, color=DARK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, *, size=10.5, bold=False, color=DARK, name="DejaVu Sans"):
    run = paragraph.add_run(text)
    set_run_font(run, name=name, size=size, bold=bold, color=color)
    return run


def add_page_title(doc, title: str, intro: str | None = None, *, break_before: bool = False) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = break_before
    p.paragraph_format.space_after = Pt(10)
    add_text(p, title, size=23, bold=True, color="000000")
    if intro:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(18)
        add_text(p, intro, size=11.5, color=SLATE)


def add_body(doc, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True)
        add_text(p, text[len(bold_lead):])
    else:
        add_text(p, text)


def add_footer(section) -> None:
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(7.2))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(left.paragraphs[0], "FACTHOR8  /  MANUAL DE MARCA", size=7.5, bold=True, color=SLATE)
    add_text(right.paragraphs[0], "VERSIÓN 1.0  /  2026", size=7.5, color=SLATE)


def add_spec_table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, label in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        add_text(cell.paragraphs[0], label, size=9, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            set_cell_fill(cell, "FFFFFF" if row_index % 2 == 0 else "F1F5F8")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            add_text(cell.paragraphs[0], value, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_logo_image(doc, filename: str, width: float, center=True) -> None:
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.add_run().add_picture(str(PNG / filename), width=Inches(width))


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "DejaVu Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "DejaVu Sans")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)

    # Cover
    doc.add_paragraph().paragraph_format.space_after = Pt(44)
    add_logo_image(doc, "facthor8-horizontal-3200.png", 7.0)
    p = doc.add_paragraph(style="Title")
    remove_paragraph_border(doc.styles["Title"]._element)
    remove_paragraph_border(p._p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(48)
    p.paragraph_format.space_after = Pt(8)
    add_text(p, "Manual de identidad visual", size=31, bold=True, color="000000")
    p = doc.add_paragraph()
    add_text(p, "Sistema de marca y normas de aplicación", size=14, color=SLATE)
    doc.add_paragraph().paragraph_format.space_after = Pt(76)
    p = doc.add_paragraph()
    add_text(p, "THE HUMAN SECURITY LAYER", size=10, bold=True, color=GREEN)
    p = doc.add_paragraph()
    add_text(p, "Versión 1.0   |   Septiembre 2026", size=9, color=SLATE)

    # Concept
    add_page_title(doc, "Concepto de marca", "Facthor8 representa la capa que ninguna tecnología puede sustituir: las decisiones humanas.", break_before=True)
    add_body(doc, "Los siete sectores azules representan las capas técnicas y organizativas que protegen una empresa. El octavo sector, en verde y desplazado hacia afuera, representa a la persona: forma parte del sistema, pero necesita atención específica.")
    add_body(doc, "El círculo verde completa una figura humana mínima. La marca evita presentar a las personas como el eslabón débil; las muestra como una capa activa que puede aprender, detectar y responder.")
    add_logo_image(doc, "facthor8-symbol-2048.png", 2.85)
    add_spec_table(doc, ["Elemento", "Significado", "Regla"], [
        ["Siete bloques azules", "Capas de protección tecnológica y organizativa", "Siempre idénticos y equidistantes"],
        ["Bloque verde", "Octava capa o factor humano", "Misma geometría, desplazada hacia afuera"],
        ["Círculo verde", "Persona y capacidad de decisión", "Nunca separar ni recolocar"],
        ["Anillo abierto", "Sistema conectado y mejora continua", "Mantener el centro despejado"],
    ], [1.55, 3.0, 2.25])

    # Construction
    add_page_title(doc, "Construcción del símbolo", "La consistencia geométrica permite reconocer la marca desde un favicon hasta una presentación corporativa.", break_before=True)
    add_logo_image(doc, "facthor8-symbol-2048.png", 2.55)
    add_spec_table(doc, ["Parámetro", "Especificación"], [
        ["Sectores", "8 sectores de 39 grados"],
        ["Distribución", "Centros separados cada 45 grados"],
        ["Separaciones", "8 huecos radiales uniformes"],
        ["Sector humano", "Misma base geométrica con desplazamiento radial"],
        ["Centro", "Área circular libre sin símbolos adicionales"],
        ["Construcción", "Formas planas y colores sólidos"],
    ], [2.0, 4.8])
    add_body(doc, "No deben modificarse individualmente los sectores ni compensarse visualmente con tamaños diferentes. La jerarquía del elemento humano se logra únicamente mediante color y desplazamiento.")

    # Versions
    add_page_title(doc, "Versiones del logo", "Utilizar siempre la versión que garantice contraste limpio y lectura inmediata.", break_before=True)
    p = doc.add_paragraph()
    add_text(p, "Versión principal para fondos claros", size=11, bold=True)
    add_logo_image(doc, "facthor8-horizontal-3200.png", 6.8)
    p = doc.add_paragraph()
    add_text(p, "Versión inversa para fondos oscuros", size=11, bold=True)
    add_logo_image(doc, "facthor8-horizontal-dark-3200.png", 6.8)
    p = doc.add_paragraph()
    add_text(p, "Versión monocromática", size=11, bold=True)
    add_logo_image(doc, "facthor8-horizontal-navy-3200.png", 6.8)
    add_body(doc, "El isotipo puede utilizarse sin el nombre únicamente cuando Facthor8 ya esté identificado por contexto, como en avatares, favicons, marcas de agua o elementos repetitivos.")

    # Clear space
    add_page_title(doc, "Área de seguridad y tamaño mínimo", "El espacio libre protege la identidad y evita que otros elementos compitan con ella.", break_before=True)
    add_logo_image(doc, "facthor8-horizontal-3200.png", 6.9)
    add_body(doc, "Área de seguridad. Mantener alrededor del logo un espacio mínimo equivalente al diámetro del círculo verde del isotipo. Ningún texto, imagen, borde o elemento interactivo debe invadir esa zona.", bold_lead="Área de seguridad.")
    add_spec_table(doc, ["Aplicación", "Logo horizontal", "Isotipo"], [
        ["Pantalla", "160 px de ancho mínimo", "24 px de ancho mínimo"],
        ["Impresión", "36 mm de ancho mínimo", "8 mm de ancho mínimo"],
        ["Cabecera web", "200 a 240 px recomendado", "Solo en navegación compacta"],
        ["Avatar social", "No recomendado", "Usar isotipo centrado"],
    ], [2.1, 2.35, 2.35])
    add_body(doc, "Por debajo de estos tamaños debe priorizarse el isotipo. La leyenda The Human Security Layer puede omitirse solo en aplicaciones compactas previamente definidas.")

    # Color
    add_page_title(doc, "Paleta cromática", "El azul construye confianza y rigor; el verde identifica exclusivamente el factor humano y las acciones prioritarias.", break_before=True)
    rows = [
        ["Midnight Navy", "#06294D", "6 41 77", "92 47 0 70", "Logo, titulares y superficies"],
        ["Deep Navy", "#041C34", "4 28 52", "92 46 0 80", "Fondos y secciones inmersivas"],
        ["Human Green", "#00D39A", "0 211 154", "100 0 27 17", "Capa humana y llamadas a la acción"],
        ["Mineral White", "#F7FAFC", "247 250 252", "2 1 0 1", "Fondos claros y texto inverso"],
        ["Signal Slate", "#7890A5", "120 144 165", "27 13 0 35", "Texto secundario y datos"],
    ]
    table = add_spec_table(doc, ["Nombre", "HEX", "RGB", "CMYK aprox.", "Uso principal"], rows, [1.25, 1.0, 1.2, 1.25, 2.1])
    swatch_colors = [NAVY, DARK, GREEN, ICE, SLATE]
    for i, color in enumerate(swatch_colors, start=1):
        set_cell_fill(table.rows[i].cells[0], color)
        text_color = "FFFFFF" if color in {NAVY, DARK, SLATE} else DARK
        for run in table.rows[i].cells[0].paragraphs[0].runs:
            run.font.color.rgb = RGBColor.from_string(text_color)
    add_body(doc, "Human Green no debe emplearse como color de texto pequeño sobre blanco. En botones verdes se utilizará texto Deep Navy para mantener un contraste funcional.")

    # Typography
    add_page_title(doc, "Tipografía", "El sistema combina expresión editorial, lectura clara y precisión técnica.", break_before=True)
    typographic_rows = [
        ["Logotipo", "Construcción propietaria", "No recrear escribiendo el nombre"],
        ["Titulares", "Syne 600 y 700", "Mensajes principales y secciones"],
        ["Texto", "Manrope 400 a 700", "Párrafos, navegación y botones"],
        ["Datos", "DM Mono 400 y 500", "Métricas, códigos y etiquetas"],
    ]
    add_spec_table(doc, ["Función", "Familia", "Aplicación"], typographic_rows, [1.45, 2.2, 3.35])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    add_text(p, "Personas preparadas. Riesgo medible.", size=24, bold=True, color="000000")
    p = doc.add_paragraph()
    add_text(p, "La seguridad se convierte en comportamiento cotidiano cuando las personas pueden practicar, decidir y aprender de forma continua.", size=12, color=SLATE)
    p = doc.add_paragraph()
    add_text(p, "HCRS / 072 / MODELO ACTIVO", size=10, bold=True, color=GREEN, name="DejaVu Sans Mono")
    add_body(doc, "No sustituir las familias principales por tipografías decorativas, condensadas o de estética hacker. Cuando no estén disponibles, utilizar Arial como alternativa operativa.")

    # Usage
    add_page_title(doc, "Usos correctos e incorrectos", "La consistencia importa más que la variedad. No deben crearse versiones nuevas para resolver aplicaciones puntuales.", break_before=True)
    add_spec_table(doc, ["Correcto", "Incorrecto"], [
        ["Usar los archivos maestros", "Redibujar o reconstruir el logo"],
        ["Mantener proporciones y área de seguridad", "Estirar, comprimir, rotar o inclinar"],
        ["Aplicar la paleta oficial", "Cambiar colores o introducir degradados"],
        ["Elegir la versión con contraste suficiente", "Añadir contornos, sombras o resplandores"],
        ["Usar el isotipo en tamaños pequeños", "Eliminar sectores o mover la figura humana"],
        ["Mantener la leyenda en piezas institucionales", "Modificar el claim o traducirlo dentro del logo"],
    ], [3.4, 3.4])
    add_body(doc, "Sobre fotografías, el logo solo podrá colocarse en una zona visualmente limpia. Si no existe contraste estable, deberá emplearse sobre una placa Deep Navy o Mineral White.")

    # Applications and package
    add_page_title(doc, "Aplicaciones y archivos", "El sistema incluye variantes preparadas para web, documentos, presentaciones, impresión y redes.", break_before=True)
    add_spec_table(doc, ["Archivo", "Uso"], [
        ["facthor8-horizontal.svg", "Logo principal sobre fondos claros"],
        ["facthor8-horizontal-reverse.svg", "Logo inverso transparente sobre fondos oscuros"],
        ["facthor8-horizontal-dark.svg", "Logo sobre placa Deep Navy"],
        ["facthor8-horizontal-navy.svg", "Aplicación monocromática oscura"],
        ["facthor8-horizontal-white.svg", "Aplicación monocromática clara"],
        ["facthor8-symbol.svg", "Isotipo, favicon, avatar y marca de agua"],
        ["PNG 3200 px", "Presentaciones y documentos de alta resolución"],
        ["PNG isotipo 2048 px", "Redes sociales y aplicaciones raster"],
    ], [3.4, 3.4])
    add_body(doc, "Web. Usar Deep Navy como fondo dominante y Human Green solo para acciones, métricas y señales importantes. El verde no debe superar aproximadamente el 15 % de la superficie visual.", bold_lead="Web.")
    add_body(doc, "Documentos. Priorizar fondos claros, logo principal y jerarquías sobrias. Reservar la versión inversa para portadas o separadores oscuros.", bold_lead="Documentos.")
    add_body(doc, "Redes. Utilizar el isotipo centrado con suficiente margen. No intentar incluir el wordmark completo dentro de avatares pequeños.", bold_lead="Redes.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
