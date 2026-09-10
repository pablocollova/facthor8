#!/usr/bin/env python3
"""Build the Facthor8 commercial dossier for Bettergy."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "brand" / "dossier"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "Facthor8_Dosier_Bettergy.docx"
LOGO = ROOT / "brand" / "png" / "facthor8-horizontal-3200.png"
SYMBOL = ROOT / "brand" / "png" / "facthor8-symbol-2048.png"

NAVY = "06294D"
DARK = "041C34"
GREEN = "00D39A"
ICE = "F7FAFC"
SLATE = "7890A5"
PALE = "EDF3F6"
BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), color)


def borders(cell, color=BORDER, size="5"):
    props = cell._tc.get_or_add_tcPr()
    node = props.first_child_found_in("w:tcBorders")
    if node is None:
        node = OxmlElement("w:tcBorders")
        props.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = node.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            node.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def cell_margins(cell, top=130, start=145, bottom=130, end=145):
    props = cell._tc.get_or_add_tcPr()
    node = props.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        props.append(node)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = node.find(qn(f"w:{name}"))
        if item is None:
            item = OxmlElement(f"w:{name}")
            node.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def run_style(run, size=10.8, bold=False, color=DARK, font="Liberation Sans"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def para_style(p, before=0, after=7, line=1.14, keep=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep


def add_text(doc, text, size=10.8, color=DARK, bold_lead=None, after=7):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        run_style(r, size=size, bold=True, color=color)
        text = text[len(bold_lead):]
    r = p.add_run(text)
    run_style(r, size=size, color=color)
    para_style(p, after=after)
    return p


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        run_style(r, size=10.4, bold=True)
        text = text[len(bold_lead):]
    r = p.add_run(text)
    run_style(r, size=10.4)
    para_style(p, after=4, line=1.1)


def add_label(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text.upper())
    run_style(r, size=8.2, bold=True, color="007B5D")
    para_style(p, after=9, keep=True)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    run_style(r, size=26 if level == 1 else 15.5, bold=True, color=BLACK, font="DejaVu Sans")
    para_style(p, before=0 if level == 1 else 10, after=12 if level == 1 else 6, line=1.0, keep=True)
    return p


def page(doc, number, label, title, intro=None):
    doc.add_page_break()
    add_label(doc, f"{number}  {label}")
    add_heading(doc, title)
    if intro:
        add_text(doc, intro, size=11.2, color=SLATE, after=13)


def table(doc, headers, rows, widths, font_size=9.0):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for i, value in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, NAVY)
        borders(c)
        cell_margins(c, 140, 140, 140, 140)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = c.paragraphs[0].add_run(value)
        run_style(r, size=9.0, bold=True, color=WHITE)
        para_style(c.paragraphs[0], after=0, line=1.05)
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(header)
    for row_idx, values in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(values):
            c = cells[i]
            shade(c, WHITE if row_idx % 2 == 0 else PALE)
            borders(c)
            cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            r = p.add_run(value)
            run_style(r, size=font_size, bold=(i == 0), color=DARK)
            para_style(p, after=0, line=1.08)
    return t


def setup(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.78)
    section.right_margin = Inches(.78)
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    normal.font.size = Pt(10.8)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "DejaVu Sans"
        style._element.rPr.rFonts.set(qn("w:ascii"), "DejaVu Sans")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "DejaVu Sans")
        style.font.color.rgb = RGBColor.from_string(BLACK)
    footer = section.footer
    ft = footer.add_table(rows=1, cols=2, width=Inches(6.9))
    ft.columns[0].width = Inches(5.4)
    ft.columns[1].width = Inches(1.5)
    left = ft.cell(0, 0).paragraphs[0]
    r = left.add_run("FACTHOR8   /   THE HUMAN SECURITY LAYER")
    run_style(r, size=7.2, bold=True, color=SLATE)
    right = ft.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = right.add_run("Página ")
    run_style(r, size=7.2, color=SLATE)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    right._p.append(field)
    for c in ft.rows[0].cells:
        cell_margins(c, 0, 0, 0, 0)


def build():
    doc = Document()
    setup(doc)

    # Cover
    p = doc.add_paragraph()
    p.add_run().add_picture(str(LOGO), width=Inches(6.55))
    para_style(p, after=24)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("08")
    run_style(r, size=78, bold=True, color=GREEN, font="DejaVu Sans")
    para_style(p, after=8)
    title = doc.add_paragraph(style="Title")
    title.add_run("Plan de resiliencia humana y protección frente al fraude digital")
    para_style(title, after=14, line=.96)
    sub = doc.add_paragraph()
    r = sub.add_run("Propuesta de intervención para Bettergy")
    run_style(r, size=16, color=SLATE)
    para_style(sub, after=24)
    add_text(doc, "Un plan coordinado entre IT y RRHH para convertir un incidente en capacidad organizacional: personas preparadas, canales claros, controles revisados y evidencia de mejora.", size=12, after=20)
    add_text(doc, "Documento de alcance inicial  /  Septiembre 2026", size=9, color="007B5D")

    page(doc, "01", "Punto de partida", "La resiliencia se construye antes del siguiente incidente", "Resolver técnicamente un incidente contiene el efecto inmediato. Convertirlo en aprendizaje reduce la probabilidad y el impacto de la próxima amenaza.")
    add_text(doc, "Los ataques actuales aprovechan cuentas legítimas, marcas conocidas, urgencia y entornos cotidianos como Microsoft 365, SharePoint y OneDrive. Por eso una organización no puede depender únicamente de filtros o herramientas: necesita que las personas sepan reconocer señales débiles, detener una acción y reportarla con rapidez.")
    add_heading(doc, "La octava capa", 2)
    add_text(doc, "Facthor8 trabaja en el punto donde los controles tecnológicos se encuentran con una decisión humana. Nuestro objetivo no es culpabilizar ni generar miedo. Es desarrollar resiliencia: la capacidad de anticipar, responder, recuperarse y aprender sin bloquear la operación.")
    add_bullet(doc, "Claridad para decidir ante mensajes y solicitudes ambiguas.", "Claridad")
    add_bullet(doc, "Práctica repetida en escenarios relevantes para el negocio.", "Práctica")
    add_bullet(doc, "Un canal inmediato y una respuesta conocida por toda la compañía.", "Un canal")
    add_bullet(doc, "Medición para demostrar evolución y priorizar refuerzos.", "Medición")

    page(doc, "02", "Respuesta", "Cada necesidad se traduce en una medida verificable", "La propuesta cubre formación, simulación, respuesta y revisión preventiva dentro de un único plan coordinado.")
    rows = [
        ("Formación obligatoria", "Sesión breve para toda la compañía y microcontenidos posteriores", "Asistencia, comprensión y evaluación posterior"),
        ("Phishing y suplantación", "Escenarios realistas con cuentas, proveedores y solicitudes creíbles", "Tasa de interacción, reporte y tiempo de reacción"),
        ("Credenciales y MFA", "Práctica sobre robo de acceso, fatiga MFA y verificación de solicitudes", "Decisiones correctas por escenario"),
        ("SharePoint y OneDrive", "Validación de enlaces, permisos, dominios, remitentes y contexto", "Reconocimiento de señales y escalado"),
        ("Gestión de incidentes", "Protocolo breve y canal único de reporte", "Tiempo hasta el primer reporte y calidad del dato"),
        ("Microsoft 365 y Defender", "Revisión de configuración, brechas y controles adicionales", "Hallazgos priorizados y plan de remediación"),
    ]
    table(doc, ["Necesidad", "Respuesta Facthor8", "Evidencia"], rows, [1.55, 3.3, 2.05], 8.5)

    page(doc, "03", "Gobernanza", "IT y RRHH lideran una intervención común", "La seguridad aporta criterio técnico; RRHH facilita alcance, comunicación y adopción; dirección legitima la prioridad.")
    table(doc, ["Responsable", "Rol principal", "Decisiones"], [
        ("Dirección", "Patrocinio y prioridad organizacional", "Objetivos, obligatoriedad y tolerancia al riesgo"),
        ("IT y Seguridad", "Controles, incidentes, escenarios y respuesta", "Canal, protocolo, configuración y remediación"),
        ("RRHH y Comunicación", "Convocatoria, segmentación y refuerzo", "Calendario, mensajes y acompañamiento"),
        ("Facthor8", "Diagnóstico, diseño, facilitación y medición", "Metodología, simulaciones, evidencia y roadmap"),
        ("Colaboradores", "Participación, decisión y reporte", "Aplicación cotidiana del protocolo"),
    ], [1.45, 3.0, 2.45], 9.0)
    add_heading(doc, "Principios operativos", 2)
    add_bullet(doc, "Aprender sin exponer ni señalar individualmente a los participantes.")
    add_bullet(doc, "Utilizar datos agregados para mejorar controles y formación.")
    add_bullet(doc, "Separar claramente las simulaciones de cualquier acción que pueda afectar sistemas reales.")
    add_bullet(doc, "Acordar previamente alcance, privacidad, responsables y reglas de escalado.")

    page(doc, "04", "Formación", "Una formación breve que prepara para actuar", "La sesión inicial está diseñada para toda la plantilla, con lenguaje claro, ejemplos del entorno real y una duración compatible con la operación.")
    table(doc, ["Bloque", "Contenido práctico", "Tiempo"], [
        ("Contexto", "Cómo operan los ataques con cuentas legítimas y marcas conocidas", "5 min"),
        ("Phishing", "Señales, suplantación, urgencia y solicitudes fuera de patrón", "10 min"),
        ("Accesos", "Robo de credenciales, MFA, fatiga MFA y verificación", "10 min"),
        ("Nube", "Enlaces y permisos en SharePoint y OneDrive", "8 min"),
        ("Respuesta", "Detener, preservar, reportar y no propagar", "10 min"),
        ("Validación", "Desafío final y compromiso de conducta", "7 min"),
    ], [1.2, 4.7, 1.0], 9.0)
    add_text(doc, "La formación se complementa con refuerzos breves por rol y por patrón detectado. El objetivo no es memorizar una lista, sino mejorar la calidad de la decisión bajo presión.", after=0)

    page(doc, "05", "Simulaciones", "Practicar antes de enfrentar una amenaza real", "Las simulaciones periódicas permiten observar conductas, reforzar el protocolo y medir si la respuesta mejora con el tiempo.")
    add_heading(doc, "Escenarios propuestos", 2)
    add_bullet(doc, "Correo de proveedor o dirección con solicitud urgente.")
    add_bullet(doc, "Enlace de SharePoint o OneDrive enviado desde una cuenta comprometida.")
    add_bullet(doc, "Página de acceso falsa y solicitud de credenciales.")
    add_bullet(doc, "Notificación o aprobación MFA inesperada.")
    add_bullet(doc, "Suplantación por mensaje, llamada o combinación de canales.")
    add_heading(doc, "Cadencia", 2)
    add_text(doc, "Proponemos una simulación de línea base, ejercicios periódicos con dificultad progresiva y refuerzo inmediato. Las personas que reportan correctamente reciben confirmación; quienes interactúan con el escenario reciben aprendizaje contextual, breve y sin culpabilización.")
    add_heading(doc, "Qué medimos", 2)
    add_text(doc, "Interacción, entrega de datos, reporte, tiempo de reporte, calidad de la información aportada y evolución por grupos de riesgo. Los resultados se presentan de forma agregada.")

    page(doc, "06", "Respuesta", "Un protocolo simple y un único canal", "Ante un correo sospechoso, la velocidad depende menos de un documento extenso que de una secuencia conocida por todos.")
    table(doc, ["Paso", "Conducta esperada", "Responsable"], [
        ("1  Detener", "No hacer clic, responder, reenviar ni aprobar MFA", "Colaborador"),
        ("2  Preservar", "Mantener el mensaje y registrar contexto sin manipular evidencias", "Colaborador"),
        ("3  Reportar", "Usar el botón o canal único definido por la compañía", "Colaborador"),
        ("4  Clasificar", "Confirmar recepción, priorizar y abrir el flujo de respuesta", "IT y Seguridad"),
        ("5  Contener", "Aplicar acciones técnicas y avisar a las áreas necesarias", "IT y Seguridad"),
        ("6  Aprender", "Comunicar hallazgos y ajustar controles o formación", "IT, RRHH y Facthor8"),
    ], [1.25, 4.1, 1.55], 8.8)
    add_text(doc, "El canal puede implementarse mediante el botón de reporte de Microsoft, una dirección monitorizada o la herramienta interna acordada. La decisión final depende del entorno y del flujo operativo de Bettergy.", size=10.1)

    page(doc, "07", "Controles", "Revisión preventiva de Microsoft 365 y Defender", "La capa humana mejora cuando el entorno técnico facilita la decisión correcta y reduce la exposición innecesaria.")
    add_heading(doc, "Ámbitos de revisión", 2)
    add_bullet(doc, "Protección contra phishing, suplantación, URLs y adjuntos.")
    add_bullet(doc, "Configuración y cobertura de MFA, acceso condicional y cuentas privilegiadas.")
    add_bullet(doc, "Compartición externa, permisos y alertas en SharePoint y OneDrive.")
    add_bullet(doc, "Mecanismo de reporte, buzones monitorizados y trazabilidad del incidente.")
    add_bullet(doc, "Políticas de Defender, registros, alertas y controles disponibles según licencias.")
    add_heading(doc, "Entregable técnico", 2)
    add_text(doc, "La revisión produce un inventario de hallazgos, riesgo, prioridad, responsable recomendado y esfuerzo estimado. La activación de cambios será realizada por el equipo autorizado de Bettergy o incluida mediante un alcance técnico expresamente acordado; el dosier no presupone acceso administrativo ni cambios automáticos.")

    page(doc, "08", "Calendario", "Primeras medidas rápidas y seis meses de consolidación", "El plan produce mejoras visibles durante el primer mes y mantiene la práctica el tiempo suficiente para medir cambio de comportamiento.")
    table(doc, ["Periodo", "Acciones", "Resultado"], [
        ("Semanas 1 y 2", "Kickoff IT y RRHH, línea base, protocolo, canal y revisión inicial de M365", "Riesgos priorizados y respuesta definida"),
        ("Semanas 3 y 4", "Formación obligatoria, comunicación y primera simulación", "Toda la plantilla activada"),
        ("Mes 2", "Refuerzo segmentado y ajustes técnicos prioritarios", "Brechas inmediatas reducidas"),
        ("Mes 3", "Segunda simulación y revisión de respuesta", "Evolución comparativa"),
        ("Meses 4 y 5", "Escenarios progresivos, microaprendizaje y ejercicio de incidentes", "Resiliencia sostenida"),
        ("Mes 6", "Medición final, informe ejecutivo y roadmap", "Evidencia y siguientes prioridades"),
    ], [1.25, 4.15, 1.5], 8.7)

    page(doc, "09", "Medición", "Resultados que dirección puede entender y utilizar", "Facthor8 combina señales de formación, simulación, reporte y controles para mostrar avance y riesgo residual.")
    table(doc, ["Indicador", "Qué responde"], [
        ("Cobertura y finalización", "¿La formación llegó a toda la compañía?"),
        ("Detección y reporte", "¿Las personas reconocen y escalan una amenaza?"),
        ("Tiempo de reporte", "¿La organización puede reaccionar antes?"),
        ("Reincidencia y evolución", "¿El comportamiento mejora entre ejercicios?"),
        ("Hallazgos técnicos", "¿Se reducen brechas preventivas prioritarias?"),
        ("Human Cyber Risk Score", "¿Cómo evoluciona la resiliencia humana de forma ejecutiva?"),
    ], [2.25, 4.65], 9.2)
    add_heading(doc, "Entregables", 2)
    add_text(doc, "Plan de trabajo, materiales de formación, protocolo de reporte, diseño de simulaciones, registro agregado de resultados, informe de revisión preventiva, dashboard de indicadores, informe ejecutivo y roadmap priorizado.")

    page(doc, "10", "Siguiente paso", "Definir el alcance operativo", "Esta propuesta responde al criterio planteado y funciona como base para una reunión breve de definición.")
    add_heading(doc, "Para preparar la propuesta económica", 2)
    add_bullet(doc, "Cantidad de personas, sedes, idiomas y modalidades de trabajo.")
    add_bullet(doc, "Licencias y configuración actual de Microsoft 365 y Defender.")
    add_bullet(doc, "Canal de reporte existente y responsables de respuesta.")
    add_bullet(doc, "Profundidad esperada de la revisión técnica y capacidad interna de remediación.")
    add_bullet(doc, "Calendario, restricciones operativas y criterios de éxito.")
    add_text(doc, "Con esa información, Facthor8 puede presentar medidas, responsables, calendario definitivo y presupuesto, diferenciando claramente el programa humano, las simulaciones y cualquier trabajo técnico adicional.", size=11.2, after=24)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(str(SYMBOL), width=Inches(1.05))
    para_style(p, after=8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("FACTHOR8")
    run_style(r, size=20, bold=True, color=NAVY, font="DejaVu Sans")
    p.add_run("\n")
    r = p.add_run("Building Human Resilience")
    run_style(r, size=10, color="007B5D")

    doc.core_properties.title = "Plan de resiliencia humana y protección frente al fraude digital"
    doc.core_properties.subject = "Propuesta Facthor8 para Bettergy"
    doc.core_properties.author = "Facthor8"
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
